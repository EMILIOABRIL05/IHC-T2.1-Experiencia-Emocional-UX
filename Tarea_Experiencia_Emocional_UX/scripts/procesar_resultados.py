#!/usr/bin/env python3
"""
Procesa los datos reales de las pruebas con usuarios y regenera los productos:

  curva_emocional/curva_emocional.png       curva del ORIGINAL: trayectorias individuales + promedio
  curva_emocional/curva_rediseno.png        curva del REDISEÑO: trayectorias individuales + promedio
  curva_emocional/curva_comparacion.png     promedio antes vs. después
  matriz/matriz_ux.xlsx                     matriz momento/emoción/evidencia/mejora + datos + gráfico
  informe/datos_informe.json                insumo para scripts/generar_informe.js

Uso (desde la carpeta Tarea_Experiencia_Emocional_UX):
  python3 scripts/procesar_resultados.py                       # usa evidencias/pruebas_usuarios/datos_prueba.csv
  python3 scripts/procesar_resultados.py --csv otro.csv
  node scripts/generar_informe.js && soffice --headless --convert-to pdf --outdir informe informe/informe_final.docx

El CSV es el que exporta el kit de prueba (evidencias/pruebas_usuarios/kit_prueba/index.html).
Si el CSV no existe, se genera la matriz con las columnas de resultados vacías y el informe
marca como PENDIENTE las secciones que dependen de los datos. Nunca se inventan valores.
"""
import argparse, csv, json, os, statistics
from collections import Counter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
P = lambda *a: os.path.join(ROOT, *a)

VERS = {"A_original": "Original (B-4)", "B_rediseno": "Rediseño (Figma)"}

# ------------------------------------------------------------------ contenido de la inspección
MOMENTOS = [
    {"id": "M1", "nombre": "Ingreso", "fase": "Antes",
     "descripcion": "El usuario llega al paso 4 «Mapa» y recibe la primera impresión de la pantalla.",
     "captura": "C1, C2",
     "evidencia": "Banner rojo «¡OFERTA ÚNICA!» y cuenta regresiva 09:59 activa en un paso de verificación; "
                  "errores ortográficos («Ofertta», «ayuda??») y tres tipografías mezcladas; título en mayúsculas sin indicar qué hacer.",
     "metrica": "Tiempo hasta la primera acción; comentarios al pensar en voz alta.",
     "mejora": "M1 · Encabezado sobrio y aviso informativo de precio reservado, sin cuenta regresiva. "
               "M3 · Título orientado a la tarea («Confirma dónde entregaremos tu pedido») con una instrucción breve.",
     "principio": "Nielsen H8 (diseño estético y minimalista) y H2 (lenguaje del usuario); diseño emocional, nivel visceral (Norman, 2004); "
                  "evitar el patrón oscuro de urgencia falsa."},
    {"id": "M2", "nombre": "Explorar el mapa", "fase": "Durante",
     "descripcion": "El usuario busca su dirección en el mapa y compara el pin con el texto de la dirección.",
     "captura": "C3",
     "evidencia": "El pin rojo está en la manzana superior derecha, mientras el marcador «123 · Calle Inventada» está en la manzana inferior izquierda: "
                  "el mapa contradice la dirección escrita. El pin es solo un contorno, sin etiqueta.",
     "metrica": "Arrastres del pin; si el pin quedó en la manzana correcta; dudas expresadas.",
     "mejora": "M4 · Pin sólido ubicado sobre el lote de la dirección, lote resaltado, etiqueta «Tu entrega: Calle Inventada 123» "
               "y estado «El pin coincide / está lejos de la dirección escrita».",
     "principio": "Nielsen H1 (visibilidad del estado del sistema); Gestalt de proximidad (etiqueta junto al pin); "
                  "reducir la brecha de evaluación entre el modelo conceptual y el modelo mental (Norman, 2013)."},
    {"id": "M3", "nombre": "Ajustar / corregir", "fase": "Durante",
     "descripcion": "El usuario intenta mover, buscar o corregir la ubicación (o decide que no hace falta).",
     "captura": "C4",
     "evidencia": "No hay buscador ni indicación de que el pin se puede arrastrar; los controles de zoom miden 32×30 px, "
                  "no tienen etiqueta y su contraste es ≈2,2:1.",
     "metrica": "Clics de zoom, uso de búsqueda, clics sin efecto, tiempo total.",
     "mejora": "M4 · Buscador de dirección y botón «Usar mi ubicación». M5 · Pista «Arrastra el pin para ajustarlo» "
               "y botones de zoom de 44×44 px con borde.",
     "principio": "Nielsen H3 (control y libertad) y H7 (flexibilidad: buscar o arrastrar); significantes (Norman, 2013); "
                  "Ley de Fitts; WCAG 2.2 · 2.5.8 (tamaño de objetivo) y 1.4.11 (contraste no textual)."},
    {"id": "M4", "nombre": "Guardar / confirmar", "fase": "Durante",
     "descripcion": "El usuario guarda o confirma la dirección de entrega.",
     "captura": "C5, C6",
     "evidencia": "«Guardar dirección» es gris sobre beige (contraste ≈2,3:1) y parece deshabilitado; al pulsarlo no ocurre nada visible. "
                  "«OMITIR »» es el botón más grande y llamativo: la jerarquía visual invita a saltarse la verificación.",
     "metrica": "Clics en guardar; si omitió sin guardar.",
     "mejora": "M7 · Un único botón primario «Confirmar ubicación y continuar» (460×60 px, contraste 6,6:1); "
               "se elimina «OMITIR» y «Volver» queda como acción secundaria.",
     "principio": "Nielsen H5 (prevención de errores); jerarquía visual y figura-fondo (Gestalt); "
                  "WCAG 2.2 · 1.4.3 (contraste mínimo 4,5:1); Ley de Fitts."},
    {"id": "M5", "nombre": "Finalizar", "fase": "Después",
     "descripcion": "El usuario pasa al pago y valora si sabe que su dirección quedó guardada.",
     "captura": "C6 + B-5",
     "evidencia": "Tras guardar u omitir no aparece ningún mensaje; el usuario llega al pago sin confirmación de que la ubicación se guardó "
                  "y el indicador de pasos no marca los pasos completados.",
     "metrica": "Si guardó la dirección; éxito de la tarea; comentarios finales.",
     "mejora": "M8 · Mensaje «Ubicación guardada correctamente» con la dirección y opción «Deshacer»; el paso 4 se marca con ✓ "
               "y el botón cambia a «Ir al pago». M2 · Indicador de pasos con estado.",
     "principio": "Nielsen H1 (retroalimentación) y H3 (deshacer); diseño emocional, nivel reflexivo: cierre y confianza; "
                  "ISO 9241-11 (satisfacción)."},
]

PROBLEMAS = [
    ("C1", "Urgencia artificial", "Banner rojo, «¡OFERTA ÚNICA!», cuenta regresiva 09:59 y «¡¡No pierdas la Ofertta!!» en un paso de verificación."),
    ("C2", "Baja credibilidad", "Errores ortográficos («Ofertta», «ayuda??») y mezcla de tipografías serif, sans y monoespaciada."),
    ("C3", "El mapa contradice la dirección", "El pin está en otra manzana que el marcador «123 · Calle Inventada»."),
    ("C4", "Sin control para corregir", "No hay buscador ni pista de arrastre; zoom de 32×30 px con contraste ≈2,2:1."),
    ("C5", "Guardar parece deshabilitado", "Texto gris sobre beige (≈2,3:1) y sin respuesta visible al pulsarlo."),
    ("C6", "Jerarquía invertida", "«OMITIR »» domina la pantalla e invita a saltar la verificación; no hay confirmación final."),
]

MEJORAS = [
    ("M1", "Encabezado sobrio y aviso de precio sin cuenta regresiva", "Reduce la presión y la desconfianza del ingreso."),
    ("M2", "Indicador de pasos con estado (✓ completados, «Paso 4 de 5»)", "Orientación: el usuario sabe dónde está y cuánto falta."),
    ("M3", "Título orientado a la tarea + instrucción", "Claridad: dice qué hacer y cómo saber que está bien."),
    ("M4", "Buscador, «Usar mi ubicación», pin alineado y arrastrable con etiqueta", "Control: corregir es posible y visible."),
    ("M5", "Zoom de 44×44 px y pista «Arrastra el pin»", "Control y accesibilidad (Fitts, WCAG 2.5.8)."),
    ("M6", "Panel con estado «El pin coincide», referencia y costos completos", "Confianza y transparencia."),
    ("M7", "Un solo botón primario «Confirmar ubicación y continuar»; sin «OMITIR»", "Prevención de errores y jerarquía clara."),
    ("M8", "Mensaje de confirmación con «Deshacer»", "Retroalimentación y cierre: el usuario sabe que se guardó."),
]

EMO_POS = {"Confianza", "Tranquilidad", "Satisfacción", "Curiosidad"}

# ------------------------------------------------------------------ contraste WCAG (cálculo reproducible)
def _lum(h):
    h = h.lstrip("#"); c = [int(h[i:i+2], 16)/255 for i in (0, 2, 4)]
    c = [v/12.92 if v <= 0.03928 else ((v+0.055)/1.055)**2.4 for v in c]
    return 0.2126*c[0] + 0.7152*c[1] + 0.0722*c[2]
def contraste(a, b):
    la, lb = sorted([_lum(a), _lum(b)], reverse=True); return round((la+0.05)/(lb+0.05), 1)

CONTRASTES = [
    # (elemento, versión, texto, fondo, medido_en)
    ("«Guardar dirección»", "Original", "#888772", "#CECFAD", "muestreo de píxeles de la captura"),
    ("Botones de zoom (+/−)", "Original", "#8E8B78", "#CFD0A8", "muestreo de píxeles de la captura"),
    ("Etiqueta «Dirección»", "Original", "#8A8B00", "#F2FF00", "muestreo de píxeles de la captura"),
    ("Botón «Confirmar ubicación»", "Rediseño", "#FFFFFF", "#1F4FD8", "colores definidos en Figma"),
    ("Texto secundario", "Rediseño", "#4A5363", "#FFFFFF", "colores definidos en Figma"),
    ("Estado «El pin coincide»", "Rediseño", "#17693E", "#E7F5EC", "colores definidos en Figma"),
    ("Borde de zoom / campos", "Rediseño", "#767F90", "#FFFFFF", "colores definidos en Figma (contraste no textual, mín. 3:1)"),
]

# ------------------------------------------------------------------ lectura de datos
def leer_csv(path):
    if not os.path.exists(path):
        return []
    with open(path, encoding="utf-8-sig") as f:
        filas = [r for r in csv.DictReader(f)]
    out = []
    for r in filas:
        try:
            r["valoracion"] = int(float(r["valoracion"]))
        except (ValueError, KeyError):
            continue
        out.append(r)
    return out

def resumir(filas):
    """Devuelve estructura con valores por versión, participante y momento."""
    res = {}
    for v in VERS:
        fv = [r for r in filas if r["version"] == v]
        if not fv:
            continue
        # si un participante tiene varias sesiones de la misma versión, se usa la última
        ses = {}
        for r in fv:
            ses.setdefault(r["participante"], {})
            ses[r["participante"]].setdefault(r["sesion"], []).append(r)
        part = {}
        for p, d in ses.items():
            ultima = sorted(d.keys())[-1]
            part[p] = {r["momento"]: r for r in d[ultima]}
        ids = [m["id"] for m in MOMENTOS]
        prom = {}
        emo = {}
        for m in ids:
            vals = [part[p][m]["valoracion"] for p in part if m in part[p]]
            prom[m] = round(statistics.mean(vals), 2) if vals else None
            e = [part[p][m]["emocion"] for p in part if m in part[p] and part[p][m]["emocion"]]
            c = Counter(e).most_common()
            emo[m] = ", ".join(f"{k} ({n})" for k, n in c) if c else ""
        meta = {}
        for p in part:
            r0 = next(iter(part[p].values()))
            meta[p] = {k: r0.get(k, "") for k in ("tiempo_total_s", "t_primera_accion_s", "arrastres_pin", "clics_zoom", "clics_guardar",
                                                   "uso_busqueda", "guardo_direccion", "omitio_sin_guardar", "pin_correcto_al_final",
                                                   "exito_tarea", "clics_sin_efecto", "notas_observador", "observador", "fecha")}
        comentarios = [(p, m, part[p][m]["comentario"]) for p in sorted(part) for m in ids if m in part[p] and part[p][m].get("comentario")]
        res[v] = {"participantes": sorted(part), "valores": {p: {m: part[p][m]["valoracion"] for m in part[p]} for p in part},
                  "promedio": prom, "emociones": emo, "meta": meta, "comentarios": comentarios}
    return res

# ------------------------------------------------------------------ gráficos
def graficar(res):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    for f in ("/usr/share/fonts/opentype/inter/Inter-Regular.otf", "/usr/share/fonts/opentype/inter/Inter-Bold.otf"):
        if os.path.exists(f): font_manager.fontManager.addfont(f)
    plt.rcParams.update({"font.family": "Inter", "font.size": 12, "axes.edgecolor": "#BDBCB6", "axes.labelcolor": "#52514E",
                         "xtick.color": "#52514E", "ytick.color": "#52514E", "figure.facecolor": "#FFFFFF", "axes.facecolor": "#FFFFFF"})
    SERIES = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#6250d6", "#e34948"]
    INK, MUT = "#0b0b0b", "#52514e"
    ids = [m["id"] for m in MOMENTOS]
    xl = [f"{m['id']}\n{m['nombre']}\n({m['fase']})" for m in MOMENTOS]
    salidas = {}

    def base(ax, titulo):
        ax.axhline(0, color="#9A9993", lw=1, zorder=0)
        ax.set_ylim(-3.4, 3.4); ax.set_yticks(range(-3, 4)); ax.set_yticklabels(["−3", "−2", "−1", "0", "+1", "+2", "+3"])
        ax.set_xticks(range(len(ids))); ax.set_xticklabels(xl, fontsize=11)
        ax.grid(axis="y", color="#ECEBE7", lw=0.8); ax.set_axisbelow(True)
        for s in ("top", "right"): ax.spines[s].set_visible(False)
        ax.set_ylabel("Intensidad emocional (−3 muy negativa · +3 muy positiva)", fontsize=11)
        ax.set_title(titulo, loc="left", fontsize=15, fontweight="bold", color=INK, pad=14)

    for v, fname in (("A_original", "curva_emocional.png"), ("B_rediseno", "curva_rediseno.png")):
        if v not in res: continue
        d = res[v]
        fig, ax = plt.subplots(figsize=(11, 6.2), dpi=160)
        base(ax, f"Curva emocional · {VERS[v]} · n = {len(d['participantes'])} participantes")
        for i, p in enumerate(d["participantes"]):
            ys = [d["valores"][p].get(m) for m in ids]
            col = SERIES[i % len(SERIES)]
            ax.plot(range(len(ids)), ys, color=col, lw=1.6, marker="o", ms=7, mec="white", mew=1.5, alpha=0.9, label=p)
        av = [d["promedio"][m] for m in ids]
        ax.plot(range(len(ids)), av, color=INK, lw=3, marker="D", ms=9, mec="white", mew=2, label="Promedio", zorder=5)
        for x, y in enumerate(av):
            ax.annotate(f"{y:+.2f}".replace(".", ",").replace("-", "−"), (x, y), textcoords="offset points", xytext=(0, 16), ha="center", fontsize=11, fontweight="bold", color=INK,
                        bbox=dict(boxstyle="round,pad=0.25", fc="white", ec="none", alpha=0.9), zorder=7)
        if v == "A_original":
            criticos = sorted(range(len(ids)), key=lambda i: av[i])[:2]
            for i in criticos:
                ax.scatter([i], [av[i]], s=600, facecolors="none", edgecolors="#B42318", linewidths=2, zorder=6)
                ax.annotate("Momento crítico", (i, av[i]), textcoords="offset points", xytext=(0, -34), ha="center", fontsize=10, color="#B42318", fontweight="bold",
                            bbox=dict(boxstyle="round,pad=0.25", fc="white", ec="#B42318", lw=1), zorder=7)
        ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.2), ncol=min(6, len(d["participantes"])+1), frameon=False, fontsize=11)
        fig.tight_layout(); fig.savefig(P("curva_emocional", fname), bbox_inches="tight"); plt.close(fig)
        salidas[v] = fname

    if "A_original" in res and "B_rediseno" in res:
        fig, ax = plt.subplots(figsize=(11, 6.2), dpi=160)
        base(ax, "Curva emocional promedio · antes vs. después del rediseño")
        for v, col, mk, lab in (("A_original", "#eb6834", "o", "Antes · original"), ("B_rediseno", "#2a78d6", "s", "Después · rediseño")):
            av = [res[v]["promedio"][m] for m in ids]
            ax.plot(range(len(ids)), av, color=col, lw=2.5, marker=mk, ms=9, mec="white", mew=2, label=f"{lab} (n = {len(res[v]['participantes'])})")
            for x, y in enumerate(av):
                ax.annotate(f"{y:+.2f}".replace(".", ",").replace("-", "−"), (x, y), textcoords="offset points", xytext=(0, 14 if v == "B_rediseno" else -22), ha="center", fontsize=10, color=MUT,
                            bbox=dict(boxstyle="round,pad=0.2", fc="white", ec="none", alpha=0.9), zorder=7)
        ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.2), ncol=2, frameon=False, fontsize=11)
        fig.tight_layout(); fig.savefig(P("curva_emocional", "curva_comparacion.png"), bbox_inches="tight"); plt.close(fig)
        salidas["comparacion"] = "curva_comparacion.png"
    return salidas

# ------------------------------------------------------------------ matriz Excel
def matriz_xlsx(res, participantes_max=6):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.chart import LineChart, Reference
    from openpyxl.utils import get_column_letter as L
    wb = Workbook()
    thin = Side(style="thin", color="D6DCE5"); B = Border(left=thin, right=thin, top=thin, bottom=thin)
    H = PatternFill("solid", fgColor="1F3A8A"); HF = Font(bold=True, color="FFFFFF")
    IN = PatternFill("solid", fgColor="FFF8DC"); wrap = Alignment(wrap_text=True, vertical="top")
    ids = [m["id"] for m in MOMENTOS]

    # --- Hoja Datos (valores reales; celdas amarillas = entrada)
    ws = wb.active; ws.title = "Datos"
    ws["A1"] = "Valoraciones emocionales por participante (escala −3 a +3). Celdas amarillas: datos reales de la prueba."; ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "Se llenan automáticamente con scripts/procesar_resultados.py o a mano. No escribir valores que no se hayan medido."
    bloques = {}
    fila = 4
    for v in VERS:
        ws.cell(fila, 1, f"Versión {VERS[v]}").font = Font(bold=True, size=12, color="1F3A8A"); fila += 1
        hdr = ["Participante"] + [f"{m['id']} {m['nombre']}" for m in MOMENTOS]
        for j, h in enumerate(hdr, 1):
            c = ws.cell(fila, j, h); c.fill = H; c.font = HF; c.border = B; c.alignment = Alignment(wrap_text=True, horizontal="center")
        fila += 1; ini = fila
        parts = res.get(v, {}).get("participantes", [])
        for i in range(max(participantes_max, len(parts))):
            p = parts[i] if i < len(parts) else ""
            ws.cell(fila, 1, p or f"P{i+1}").border = B
            for j, m in enumerate(ids, 2):
                val = res[v]["valores"][p].get(m) if p else None
                c = ws.cell(fila, j, val); c.fill = IN; c.border = B; c.alignment = Alignment(horizontal="center")
            fila += 1
        fin = fila - 1
        ws.cell(fila, 1, "Promedio").font = Font(bold=True)
        for j in range(2, 2+len(ids)):
            c = ws.cell(fila, j, f'=IFERROR(ROUND(AVERAGE({L(j)}{ini}:{L(j)}{fin}),2),"")'); c.font = Font(bold=True); c.border = B; c.alignment = Alignment(horizontal="center")
        bloques[v] = {"ini": ini, "fin": fin, "prom": fila}
        fila += 3
    ws.column_dimensions["A"].width = 16
    for j in range(2, 2+len(ids)): ws.column_dimensions[L(j)].width = 18

    # gráfico de promedios (se actualiza solo al escribir datos)
    ch = LineChart(); ch.title = "Curva emocional promedio"; ch.y_axis.title = "−3 … +3"; ch.y_axis.scaling.min = -3; ch.y_axis.scaling.max = 3
    ch.height = 9; ch.width = 20
    for v in VERS:
        r = bloques[v]["prom"]
        ref = Reference(ws, min_col=2, max_col=1+len(ids), min_row=r)
        ch.add_data(ref, from_rows=True, titles_from_data=False)
    ch.series[0].tx = None
    from openpyxl.chart.series import SeriesLabel
    ch.series[0].title = SeriesLabel(v="Antes · original"); ch.series[1].title = SeriesLabel(v="Después · rediseño")
    ch.series[0].graphicalProperties.line.solidFill = "EB6834"; ch.series[1].graphicalProperties.line.solidFill = "2A78D6"
    for s_ in ch.series:
        s_.smooth = False; s_.marker.symbol = "circle"; s_.marker.size = 7; s_.graphicalProperties.line.width = 28000
    ch.set_categories(Reference(ws, min_col=2, max_col=1+len(ids), min_row=bloques["A_original"]["ini"]-1))
    ws.add_chart(ch, f"I4")

    # --- Hoja Matriz UX
    mx = wb.create_sheet("Matriz UX", 0)
    cols = ["Momento", "Fase", "Captura (ID)", "Emoción reportada (original)", "Valoración promedio ANTES",
            "Evidencia observable en la interfaz", "Evidencia registrada en la prueba", "Mejora propuesta (rediseño)",
            "Principio IHC que la fundamenta", "Emoción reportada (rediseño)", "Valoración promedio DESPUÉS", "Variación (Δ)", "Resultado de la validación"]
    widths = [20, 10, 12, 24, 13, 48, 36, 46, 44, 24, 13, 11, 26]
    mx["A1"] = "Matriz UX · momento / emoción / evidencia / mejora · Pantalla «Mapa» (B-4) · SuperOfertas Ecuador"; mx["A1"].font = Font(bold=True, size=13)
    mx["A2"] = "Escala −3 (muy negativa) a +3 (muy positiva). Valoraciones y emociones = datos reales de la prueba; se calculan desde la hoja «Datos»."
    for j, (h, w) in enumerate(zip(cols, widths), 1):
        c = mx.cell(4, j, h); c.fill = H; c.font = HF; c.border = B; c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        mx.column_dimensions[L(j)].width = w
    mx.row_dimensions[4].height = 45
    for i, m in enumerate(MOMENTOS):
        r = 5 + i; jcol = L(2 + i)
        a, b = bloques["A_original"]["prom"], bloques["B_rediseno"]["prom"]
        reg = []
        for v in VERS:
            meta = res.get(v, {}).get("meta", {})
            if not meta: continue
        valores = [f"{m['id']} · {m['nombre']}", m["fase"], m["captura"],
                   res.get("A_original", {}).get("emociones", {}).get(m["id"], ""),
                   f"=Datos!{jcol}{a}", m["evidencia"], m["metrica"], m["mejora"], m["principio"],
                   res.get("B_rediseno", {}).get("emociones", {}).get(m["id"], ""),
                   f"=Datos!{jcol}{b}", f'=IF(OR(E{r}="",K{r}=""),"",ROUND(K{r}-E{r},2))',
                   f'=IF(L{r}="","Pendiente de prueba",IF(L{r}>=1,"Mejora clara",IF(L{r}>0,"Mejora leve",IF(L{r}=0,"Sin cambio","Empeora"))))']
        for j, val in enumerate(valores, 1):
            c = mx.cell(r, j, val); c.border = B; c.alignment = wrap
            if j in (4, 10) and not val: c.fill = IN
        for j in (5, 11, 12): mx.cell(r, j).alignment = Alignment(horizontal="center", vertical="top"); mx.cell(r, j).font = Font(bold=True)
        mx.row_dimensions[r].height = 150
    mx.freeze_panes = "B5"

    # --- Hoja Métricas
    mt = wb.create_sheet("Métricas")
    hdr = ["Versión", "Participante", "Tiempo total (s)", "1.ª acción (s)", "Arrastres pin", "Clics zoom", "Clics guardar", "Usó búsqueda",
           "Guardó dirección", "Omitió sin guardar", "Pin correcto al final", "Éxito tarea", "Clics sin efecto", "Notas del observador"]
    for j, h in enumerate(hdr, 1):
        c = mt.cell(1, j, h); c.fill = H; c.font = HF; c.border = B; c.alignment = Alignment(wrap_text=True)
        mt.column_dimensions[L(j)].width = 14 if j < 14 else 60
    r = 2
    keys = ["tiempo_total_s", "t_primera_accion_s", "arrastres_pin", "clics_zoom", "clics_guardar", "uso_busqueda", "guardo_direccion",
            "omitio_sin_guardar", "pin_correcto_al_final", "exito_tarea", "clics_sin_efecto", "notas_observador"]
    for v in VERS:
        for p, meta in sorted(res.get(v, {}).get("meta", {}).items()):
            vals = [VERS[v], p] + [meta.get(k, "") for k in keys]
            for j, val in enumerate(vals, 1):
                try: val = float(val) if j in (3, 4, 5, 6, 7, 13) and val != "" else val
                except ValueError: pass
                c = mt.cell(r, j, val); c.border = B; c.alignment = wrap
            r += 1
    if r == 2:
        mt.cell(2, 1, "Pendiente: se completa con el CSV exportado por el kit de prueba.")
    for sh in wb.worksheets:
        sh.page_setup.orientation = "landscape"; sh.page_setup.paperSize = sh.PAPERSIZE_A4
        sh.sheet_properties.pageSetUpPr.fitToPage = True; sh.page_setup.fitToWidth = 1; sh.page_setup.fitToHeight = 0
    wb.save(P("matriz", "matriz_ux.xlsx"))

# ------------------------------------------------------------------ main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--csv", default=P("evidencias", "pruebas_usuarios", "datos_prueba.csv"))
    ap.add_argument("--repo", default="", help="URL del repositorio GitHub del grupo")
    ap.add_argument("--figma", default="https://www.figma.com/design/qHv7b4pqeDrXumpQNNd2FN")
    a = ap.parse_args()
    filas = leer_csv(a.csv)
    res = resumir(filas)
    graficos = graficar(res) if res else {}
    matriz_xlsx(res)
    datos = {"momentos": MOMENTOS, "problemas": PROBLEMAS, "mejoras": MEJORAS,
             "contrastes": [{"elemento": e, "version": v, "texto": t, "fondo": f, "ratio": contraste(t, f), "fuente": s} for e, v, t, f, s in CONTRASTES],
             "resultados": res, "graficos": graficos, "hay_datos": bool(res), "hay_comparacion": "comparacion" in graficos,
             "figma": a.figma, "repo": a.repo, "csv": os.path.relpath(a.csv, ROOT) if os.path.exists(a.csv) else ""}
    with open(P("informe", "datos_informe.json"), "w", encoding="utf-8") as f:
        json.dump(datos, f, ensure_ascii=False, indent=1)
    n = {v: len(res[v]["participantes"]) for v in res}
    print("Participantes por versión:", n or "sin datos (CSV no encontrado)")
    for v in res:
        print(VERS[v], "promedios:", res[v]["promedio"])
    print("Gráficos:", graficos or "ninguno (faltan datos)")
    print("Matriz: matriz/matriz_ux.xlsx · JSON: informe/datos_informe.json")

if __name__ == "__main__":
    main()
