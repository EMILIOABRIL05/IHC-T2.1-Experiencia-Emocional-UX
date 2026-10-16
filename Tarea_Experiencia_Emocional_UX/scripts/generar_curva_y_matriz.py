#!/usr/bin/env python3
"""Genera la curva emocional ESPERADA (hipótesis de la inspección) y la matriz UX.

Salidas:
  curva_emocional/curva_emocional.png   curva esperada original vs. rediseño (hipótesis, no medida)
  matriz/matriz_ux.xlsx                 matriz momento / emoción / evidencia / mejora / principio / verificación

Uso:  python3 scripts/generar_curva_y_matriz.py
"""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from contenido import MOMENTOS, VERIFICACION, HEURISTICAS, CONTRASTES, contraste

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
P = lambda *a: os.path.join(ROOT, *a)


def curva():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    for f in ("/usr/share/fonts/opentype/inter/Inter-Regular.otf", "/usr/share/fonts/opentype/inter/Inter-Bold.otf"):
        if os.path.exists(f):
            font_manager.fontManager.addfont(f)
            plt.rcParams["font.family"] = "Inter"
    plt.rcParams.update({"font.size": 12, "axes.edgecolor": "#BDBCB6", "xtick.color": "#52514E", "ytick.color": "#52514E"})
    INK, MUT = "#0b0b0b", "#52514e"
    ids = [m["id"] for m in MOMENTOS]
    xl = [f"{m['id']}\n{m['nombre']}\n({m['fase']})" for m in MOMENTOS]
    a = [m["esperado_antes"] for m in MOMENTOS]
    b = [m["esperado_despues"] for m in MOMENTOS]
    fig, ax = plt.subplots(figsize=(11, 6.4), dpi=170)
    ax.axhspan(-3.9, 0, color="#FBEDEA", zorder=0); ax.axhspan(0, 3.4, color="#EEF3FB", zorder=0)
    ax.axhline(0, color="#9A9993", lw=1, zorder=1)
    ax.set_ylim(-3.9, 3.4); ax.set_xlim(-0.45, 4.45); ax.set_yticks(range(-3, 4)); ax.set_yticklabels(["−3", "−2", "−1", "0", "+1", "+2", "+3"])
    ax.set_xticks(range(5)); ax.set_xticklabels(xl, fontsize=11)
    ax.grid(axis="y", color="#E4E3DF", lw=0.8); ax.set_axisbelow(True)
    for s in ("top", "right"): ax.spines[s].set_visible(False)
    ax.set_ylabel("Intensidad emocional esperada\n(−3 muy negativa · +3 muy positiva)", fontsize=11, color=MUT)
    ax.plot(range(5), a, color="#eb6834", lw=2.6, ls="--", marker="o", ms=10, mec="white", mew=2, label="Original B-4 (hipótesis)", zorder=4)
    ax.plot(range(5), b, color="#2a78d6", lw=2.6, ls="--", marker="s", ms=10, mec="white", mew=2, label="Rediseño Figma (hipótesis)", zorder=4)
    for x, (ya, yb, m) in enumerate(zip(a, b, MOMENTOS)):
        ax.annotate(f"{ya:+d}".replace("-", "−") + f"  {m['emocion'].split(' /')[0]}", (x, ya), textcoords="offset points", xytext=(0, -24),
                    ha="center", fontsize=10, color=INK, bbox=dict(boxstyle="round,pad=0.25", fc="white", ec="none", alpha=0.95), zorder=6)
        ax.annotate(f"{yb:+d}  {m['emocion_despues']}", (x, yb), textcoords="offset points", xytext=(0, 15),
                    ha="center", fontsize=10, color=INK, bbox=dict(boxstyle="round,pad=0.25", fc="white", ec="none", alpha=0.95), zorder=6)
    crit = [i for i, v in enumerate(a) if v <= -2]
    for i in crit:
        ax.scatter([i], [a[i]], s=700, facecolors="none", edgecolors="#B42318", linewidths=2, zorder=5)
    ax.set_title("Curva emocional esperada · pantalla «Mapa» (B-4)", loc="left", fontsize=15, fontweight="bold", color=INK, pad=26)
    ax.text(0, 1.02, "HIPÓTESIS del equipo a partir de la inspección; no son valores medidos con usuarios. Círculo rojo = momento crítico esperado.",
            transform=ax.transAxes, fontsize=10, color="#B42318")
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.22), ncol=2, frameon=False, fontsize=11)
    fig.tight_layout(); fig.savefig(P("curva_emocional", "curva_emocional.png"), bbox_inches="tight"); plt.close(fig)


def matriz():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L
    wb = Workbook(); ws = wb.active; ws.title = "Matriz UX"
    thin = Side(style="thin", color="C9CFD8"); B = Border(left=thin, right=thin, top=thin, bottom=thin)
    H = PatternFill("solid", fgColor="1F3A8A"); HF = Font(bold=True, color="FFFFFF"); wrap = Alignment(wrap_text=True, vertical="top")
    ws["A1"] = "Matriz UX · momento / emoción / evidencia / mejora · Pantalla «Mapa» (B-4) · SuperOfertas Ecuador"; ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("Emoción y valoración ESPERADAS (hipótesis del equipo a partir de la inspección; escala −3 a +3). "
                "No se realizaron pruebas con usuarios: estos valores deben validarse con participantes reales.")
    ws["A2"].font = Font(italic=True, color="B42318")
    cols = ["Momento", "Fase", "Captura", "Emoción esperada (original)", "Valor esperado antes", "Evidencia observable en la interfaz",
            "Mejora propuesta", "Principio IHC que la fundamenta", "Emoción esperada (rediseño)", "Valor esperado después"]
    widths = [20, 10, 10, 20, 11, 50, 46, 46, 18, 11]
    for j, (h, w) in enumerate(zip(cols, widths), 1):
        c = ws.cell(4, j, h); c.fill = H; c.font = HF; c.border = B; c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        ws.column_dimensions[L(j)].width = w
    ws.row_dimensions[4].height = 42
    for i, m in enumerate(MOMENTOS):
        r = 5 + i
        vals = [f"{m['id']} · {m['nombre']}", m["fase"], m["captura"], m["emocion"], m["esperado_antes"], m["evidencia"], m["mejora"],
                m["principio"], m["emocion_despues"], m["esperado_despues"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(r, j, v); c.border = B; c.alignment = wrap
        for j in (5, 10):
            ws.cell(r, j).alignment = Alignment(horizontal="center", vertical="top"); ws.cell(r, j).font = Font(bold=True)
        ws.row_dimensions[r].height = 110
    r = 5 + len(MOMENTOS)
    ws.cell(r, 4, "Promedio").font = Font(bold=True)
    ws.cell(r, 5, f"=AVERAGE(E5:E{r-1})").font = Font(bold=True); ws.cell(r, 10, f"=AVERAGE(J5:J{r-1})").font = Font(bold=True)
    ws.freeze_panes = "B5"

    v = wb.create_sheet("Verificación")
    v["A1"] = "Verificación del rediseño por inspección (valores medidos sobre los diseños)"; v["A1"].font = Font(bold=True, size=13)
    for j, (h, w) in enumerate(zip(["Criterio", "Original", "Rediseño", "Referencia"], [46, 22, 36, 26]), 1):
        c = v.cell(3, j, h); c.fill = H; c.font = HF; c.border = B; v.column_dimensions[L(j)].width = w
    for i, row in enumerate(VERIFICACION):
        for j, val in enumerate(row, 1):
            c = v.cell(4 + i, j, val); c.border = B; c.alignment = wrap
    r0 = 6 + len(VERIFICACION)
    v.cell(r0, 1, "Contraste WCAG calculado").font = Font(bold=True, size=12)
    for j, h in enumerate(["Elemento", "Versión", "Texto", "Fondo", "Contraste", "Requisito"], 1):
        c = v.cell(r0 + 1, j, h); c.fill = H; c.font = HF; c.border = B
    for i, (e, ver, t, f, req) in enumerate(CONTRASTES):
        for j, val in enumerate([e, ver, t, f, f"{contraste(t, f)}:1".replace(".", ","), req], 1):
            c = v.cell(r0 + 2 + i, j, val); c.border = B
    v.column_dimensions["E"].width = 12; v.column_dimensions["F"].width = 18

    h = wb.create_sheet("Heurísticas")
    for j, (t, w) in enumerate(zip(["Heurística (Nielsen)", "Hallazgo en el original", "Solución en el rediseño"], [32, 55, 55]), 1):
        c = h.cell(1, j, t); c.fill = H; c.font = HF; c.border = B; h.column_dimensions[L(j)].width = w
    for i, row in enumerate(HEURISTICAS):
        for j, val in enumerate(row, 1):
            c = h.cell(2 + i, j, val); c.border = B; c.alignment = wrap
    for sh in wb.worksheets:
        sh.page_setup.orientation = "landscape"; sh.page_setup.paperSize = sh.PAPERSIZE_A4
        sh.sheet_properties.pageSetUpPr.fitToPage = True; sh.page_setup.fitToWidth = 1; sh.page_setup.fitToHeight = 0
    wb.save(P("matriz", "matriz_ux.xlsx"))


if __name__ == "__main__":
    os.makedirs(P("curva_emocional"), exist_ok=True); os.makedirs(P("matriz"), exist_ok=True)
    curva(); matriz()
    print("OK: curva_emocional/curva_emocional.png · matriz/matriz_ux.xlsx")
