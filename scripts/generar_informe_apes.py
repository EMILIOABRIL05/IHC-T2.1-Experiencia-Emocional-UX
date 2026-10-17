#!/usr/bin/env python3
"""Llena la plantilla institucional FORMATO_APES.dotx con el informe de la tarea T2.1.

Uso (desde Tarea_Experiencia_Emocional_UX):
  python3 scripts/generar_curva_y_matriz.py
  python3 scripts/generar_informe_apes.py --repo https://github.com/<usuario>/<repo>
  soffice --headless --convert-to pdf --outdir informe informe/informe_final.docx
"""
import argparse, copy, io, os, sys, zipfile
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from contenido import (FIGMA, ORIGINAL_URL, TAREA, INTEGRANTES, MOMENTOS, PROBLEMAS, MEJORAS, VERIFICACION,
                       HEURISTICAS, REFERENCIAS)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
P = lambda *a: os.path.join(ROOT, *a)
W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
FONT = "Times New Roman"


# ------------------------------------------------------------------ carga de la plantilla
def abrir_plantilla(path):
    buf = io.BytesIO()
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for i in zin.infolist():
            data = zin.read(i.filename)
            if i.filename == "[Content_Types].xml":
                data = data.replace(b"template.main+xml", b"document.main+xml")
            zout.writestr(i, data)
    buf.seek(0)
    return docx.Document(buf)


class Builder:
    def __init__(self, d):
        self.d = d
        self.cursor = None
        self.nfig = 0
        self.ntab = 0

    # ---- utilidades de posición
    def at(self, el):
        self.cursor = el
        return self

    def _place(self, el):
        self.cursor.addnext(el)
        self.cursor = el
        return el

    # ---- texto
    def _runs(self, p, parts, size=None, color=None):
        if isinstance(parts, str):
            parts = [parts]
        for part in parts:
            bold = italic = False
            if isinstance(part, tuple):
                text, fmt = part
                bold, italic = "b" in fmt, "i" in fmt
            else:
                text = part
            r = p.add_run(text)
            r.font.name = FONT
            r.bold = bold or None
            r.italic = italic or None
            if size: r.font.size = Pt(size)
            if color: r.font.color.rgb = RGBColor.from_string(color)
        return p

    def par(self, parts, style="NormalAPE", align=None, after=6, before=0, size=None, keep=False, indent=None, color=None):
        p = self.d.add_paragraph(style=style)
        self._runs(p, parts, size=size, color=color)
        pf = p.paragraph_format
        pf.space_after = Pt(after); pf.space_before = Pt(before)
        if align == "c": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "l": p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if keep: pf.keep_with_next = True
        if indent is not None: pf.left_indent = Cm(indent)
        self._place(p._p)
        return p

    def sub(self, text):
        return self.par([(text, "b")], after=4, before=8, keep=True)

    def bullet(self, parts, after=3, left=False):
        p = self.d.add_paragraph(style="List Paragraph")
        pPr = p._p.get_or_add_pPr()
        numPr = OxmlElement("w:numPr")
        il = OxmlElement("w:ilvl"); il.set(qn("w:val"), "0")
        ni = OxmlElement("w:numId"); ni.set(qn("w:val"), "4")
        numPr.append(il); numPr.append(ni); pPr.append(numPr)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if left else WD_ALIGN_PARAGRAPH.JUSTIFY
        pf = p.paragraph_format; pf.space_after = Pt(after); pf.line_spacing = 1.0
        self._runs(p, parts)
        self._place(p._p)
        return p

    def blank(self):
        return self.par("", after=0)

    def page_break(self):
        p = self.d.add_paragraph(style="NormalAPE")
        p.add_run().add_break(WD_BREAK.PAGE)
        self._place(p._p)

    # ---- figuras y tablas
    def figura(self, path, ancho_cm, titulo, fuente="Elaboración propia."):
        self.nfig += 1
        p = self.d.add_paragraph(style="NormalAPE")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.left_indent = Cm(0)
        p.add_run().add_picture(P(path), width=Cm(ancho_cm))
        self._place(p._p)
        self.par([(f"Figura {self.nfig}. ", "b"), (titulo, "i")], align="c", after=0, size=10, keep=True, indent=0)
        self.par(f"Fuente: {fuente}", align="c", after=10, size=9, indent=0, color="404040")

    def figura_anexo(self, path, ancho_cm, rotulo, titulo):
        p = self.d.add_paragraph(style="NormalAPE")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.left_indent = Cm(0)
        p.add_run().add_picture(P(path), width=Cm(ancho_cm))
        self._place(p._p)
        self.par([(f"{rotulo}. ", "b"), (titulo, "i")], align="c", after=10, size=10, indent=0)

    def tabla(self, titulo, headers, rows, widths_cm, size=9, center_cols=(), fuente="Elaboración propia."):
        self.ntab += 1
        self.par([(f"Tabla {self.ntab}. ", "b"), (titulo, "i")], align="l", after=3, size=10, keep=True, before=4)
        t = self.d.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        def fill(cell, text, head=False, center=False):
            cell.text = ""
            p = cell.paragraphs[0]
            lines = str(text).split("\n")
            for k, line in enumerate(lines):
                if k: p = cell.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                if head or center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(line); r.font.name = FONT; r.font.size = Pt(size); r.bold = head or None
            if head:
                tcPr = cell._tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "D9E2F3")
                tcPr.append(shd)
        for j, h in enumerate(headers):
            fill(t.rows[0].cells[j], h, head=True)
        # repetir encabezado
        trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
        for row in rows:
            cells = t.add_row().cells
            for j, v in enumerate(row):
                fill(cells[j], v, center=j in center_cols)
        tblPr = t._tbl.tblPr
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
        for j, w in enumerate(widths_cm):
            t.columns[j].width = Cm(w)
        for row in t.rows:
            trPr = row._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trPr.append(cs)
            for j, w in enumerate(widths_cm):
                row.cells[j].width = Cm(w)
        self._place(t._tbl)
        self.par(f"Fuente: {fuente}", align="l", after=10, size=9, color="404040", before=2)
        return t


# ------------------------------------------------------------------ edición de la plantilla
def set_text(p, text, keep_first=True):
    """Deja el primer run con su formato y reemplaza el texto del párrafo."""
    runs = p.runs
    if not runs:
        r = p.add_run(text); r.font.name = FONT; return
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def set_value(p, value):
    """Para líneas «Etiqueta:<tabs>Valor»: reemplaza todo lo que va después del último tabulador."""
    runs = p.runs
    last_tab = max(i for i, r in enumerate(runs) if r._r.find(qn("w:tab")) is not None)
    after = runs[last_tab + 1:]
    if not after:
        nuevo = copy.deepcopy(runs[last_tab]._r)
        for t in nuevo.findall(qn("w:tab")): nuevo.remove(t)
        runs[last_tab]._r.addnext(nuevo)
        after = [docx.text.run.Run(nuevo, p)]
    after[0].text = value
    for r in after[1:]:
        r._r.getparent().remove(r._r)


def marcar_checkbox(p, marcado=True):
    for sdt in p._p.iter(qn("w:sdt")):
        chk = sdt.find(".//{%s}checked" % W14)
        if chk is not None:
            chk.set("{%s}val" % W14, "1" if marcado else "0")
        for t in sdt.iter(qn("w:t")):
            t.text = "☒" if marcado else "☐"


def borrar_entre(a, b):
    el = a.getnext()
    while el is not None and el is not b:
        nxt = el.getnext()
        if el.tag in (qn("w:p"), qn("w:tbl")):
            el.getparent().remove(el)
        el = nxt


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--plantilla", default=P("informe", "plantilla", "FORMATO_APES.dotx"))
    ap.add_argument("--repo", default="")
    ap.add_argument("--salida", default=P("informe", "informe_final.docx"))
    a = ap.parse_args()

    d = abrir_plantilla(a.plantilla)
    paras = d.paragraphs
    def find(text, start=0):
        for i, p in enumerate(paras[start:], start):
            if p.text.strip().startswith(text):
                return i, p
        raise KeyError(text)

    # ---------------- portada
    _, p = find("Tema"); set_value(p, "Evaluación y rediseño de la experiencia emocional del usuario")
    p.paragraph_format.left_indent = Cm(6.25); p.paragraph_format.first_line_indent = Cm(-6.25)
    tabs = [r for r in p.runs if r._r.find(qn("w:tab")) is not None]
    for r in tabs[1:]: r._r.getparent().remove(r._r)
    _, p = find("Alumnos participantes"); set_value(p, INTEGRANTES[0][0])
    prev = p._p
    for nombre in [x[0] for x in INTEGRANTES[1:]]:
        nuevo = copy.deepcopy(p._p)
        prev.addnext(nuevo); prev = nuevo
        q = docx.text.paragraph.Paragraph(nuevo, p._parent)
        q.runs[0].text = ""
        set_value(q, nombre)
        # alinear con la columna de valores
        for _ in range(2):
            r = q.runs[1]._r; r.addnext(copy.deepcopy(r))
    _, p = find("Asignatura"); set_value(p, "Interacción Humano-Computador")
    paras = d.paragraphs

    # ---------------- 2.1 Objetivos
    _, p = find("(Especificado"); set_text(p, "Analizar la experiencia emocional del usuario en la pantalla «Mapa» (B-4) del prototipo de compra "
                                              "de SuperOfertas Ecuador y proponer un rediseño en Figma fundamentado en principios de IHC, usabilidad, "
                                              "accesibilidad y diseño emocional.")
    i, p = find("(Se recomiendan")
    objetivos = [
        "Identificar cinco momentos de interacción (antes, durante y después) y los elementos de la interfaz que pueden generar confianza, incertidumbre, frustración o satisfacción.",
        "Construir la curva emocional y la matriz UX que relacione cada momento con su emoción, evidencia, mejora y principio de IHC.",
        "Rediseñar la pantalla en Figma y verificar las mejoras con criterios medibles de contraste, tamaño de objetivos, retroalimentación y heurísticas.",
    ]
    set_text(p, objetivos[0])
    for k, extra in enumerate(objetivos[1:], 1):
        q = paras[i + k]
        if q.runs: set_text(q, extra)
        else:
            r = q.add_run(extra); r.font.name = FONT; r.bold = False
    for k in (1, 2):
        for r in paras[i + k].runs: r.bold = False
    paras = d.paragraphs

    # ---------------- 2.2 / 2.3 / 2.4
    _, p = find("Detallada en la guía"); set_text(p, "Grupal (cinco integrantes).")
    _, p = find("Presenciales:"); p.runs[-1].text = ""; p.runs[1].text = "120 minutos (tiempo sugerido en la guía)."
    for r in p.runs[2:]: r._r.getparent().remove(r._r)
    _, p = find("No presenciales:")
    p.runs[-1].text = "trabajo autónomo del grupo hasta la fecha de entrega (16 de octubre de 2026)."
    paras = d.paragraphs

    b = Builder(d)
    i_ins, p_ins = find("Detalladas en la guía proporcionada")
    set_text(p_ins, "La guía de la tarea grupal T2.1 establece las siguientes actividades obligatorias:")
    b.at(p_ins._p)
    for t in ["Identificar cinco momentos de interacción: antes, durante y después.",
              "Evaluar las emociones experimentadas por al menos tres usuarios y elaborar una curva emocional.",
              "Construir una matriz que relacione momento, emoción, evidencia y propuesta de mejora.",
              "Rediseñar la interfaz en Figma y presentar una comparación visual antes/después.",
              "Verificar las mejoras mediante una nueva prueba de usuario.",
              "Publicar las evidencias y productos en GitHub con participación comprobable de los integrantes."]:
        b.bullet(t)
    paras = d.paragraphs

    # ---------------- 2.5 Equipos
    i, p = find("Listado de equipos y materiales generales")
    vacio = paras[i + 1]
    items = ["Computadoras personales de los integrantes con navegador web.",
             "Figma: recreación del original, rediseño y prototipo navegable.",
             "GitHub: repositorio del grupo y control de versiones.",
             "Microsoft Word: elaboración del informe.",
             "Python 3 (matplotlib, openpyxl): curva emocional, matriz UX y cálculo de contraste WCAG."]
    r = vacio.add_run(items[0]); r.font.name = FONT; r.bold = False
    b.at(vacio._p)
    for t in items[1:]:
        b.bullet(t, after=0)
    paras = d.paragraphs
    for label, val in (("Plataformas educativas", True), ("Aplicaciones educativas", False), ("Inteligencia Artificial", True)):
        _, p = find(label); marcar_checkbox(p, val)
    _, p = find("Otros (Especifique)"); set_text(p, "Otros (Especifique): Figma (diseño de interfaces) y GitHub (repositorio).")

    # ---------------- 2.6 Actividades por desarrollar
    paras = d.paragraphs
    i_act, h_act = find("Actividades por desarrollar")
    i_res, h_res = find("Resultados obtenidos")
    borrar_entre(h_act._p, h_res._p)
    b.at(h_act._p)
    b.par("", after=0)
    b.par("En esta sección se describe el procedimiento seguido para evaluar la experiencia emocional del usuario en la pantalla «Mapa» (B-4) "
          "del prototipo de compra de SuperOfertas Ecuador y para rediseñarla. El prototipo original está publicado en el repositorio del Grupo 1 [8].")

    b.sub("1. Caso de estudio y tarea evaluada")
    b.par("La pantalla B-4 es el cuarto de cinco pasos del proceso de compra (Carrito → Domicilio → Quién recibe → Mapa → Pagar). "
          "Su propósito es que el cliente revise en un mapa la ubicación de entrega antes de pagar. Se definió la siguiente tarea, que se usa "
          "igual para el original y para el rediseño, de modo que la comparación sea válida:")
    b.par([("Tarea: ", "b"), (f"«{TAREA}»", "i")], indent=1.2)
    b.figura("evidencias/original/B-4_anotado_C1-C6.png", 14.5, "Pantalla original B-4 con los problemas observados numerados (C1–C6).",
             "Grupo 1 HCI-2026 [8]; anotaciones del equipo.")

    b.sub("2. Actividad 1 · Inspección y momentos de interacción")
    b.par("Se inspeccionó la pantalla y su contexto (pasos B-3 y B-5) con las heurísticas de Nielsen [3], las pautas WCAG 2.2 [4], las leyes de "
          "Gestalt y la Ley de Fitts [6]. El contraste se calculó con la fórmula de luminancia relativa de WCAG, a partir de los colores "
          "muestreados en la captura original y de los colores definidos en Figma para el rediseño. La Tabla 1 presenta los cinco momentos "
          "de interacción y la Tabla 2 los problemas observados, cada uno con su captura numerada (Anexo A).")
    b.tabla("Momentos de interacción definidos para la tarea", ["Momento", "Fase", "Qué hace el usuario", "Captura"],
            [(f"{m['id']} · {m['nombre']}", m["fase"], m["que_hace"], m["captura"]) for m in MOMENTOS], [3.6, 1.8, 7.4, 2.0], center_cols=(1, 3))
    b.tabla("Problemas observados en la interfaz original", ["ID", "Problema", "Evidencia observable"],
            PROBLEMAS, [1.1, 3.8, 9.9], center_cols=(0,))
    b.par("El hallazgo más grave es C3: la pantalla pide revisar la dirección, pero el propio mapa la contradice, porque el pin está en una "
          "manzana distinta a la del número 123 de la Calle Inventada. Además, no ofrece un medio visible para corregirla (C4), y la jerarquía "
          "visual favorece «OMITIR» frente a «Guardar dirección» (C5, C6), por lo que el camino más fácil es saltarse la verificación.")

    b.sub("3. Actividad 2 · Curva emocional")
    b.par([("Alcance: ", "b"), "el grupo no realizó la prueba con usuarios ajenos al equipo. Por ese motivo, la curva que se presenta es una ",
           ("curva emocional esperada", "b"), ": cada valor es una hipótesis del equipo derivada de la evidencia observada en la inspección "
           "(Tabla 2), no una medición con participantes. Se usa la misma escala de la guía (−3 muy negativa a +3 muy positiva) y los valores "
           "deben confirmarse con usuarios reales (ver Recomendaciones)."])
    b.figura("curva_emocional/curva_emocional.png", 15.0, "Curva emocional esperada del original y del rediseño (hipótesis, no medida).")
    b.par("Según la inspección, la mayor caída se espera en M3 · Ajustar/corregir (−3, frustración): el usuario descubre que el pin no "
          "coincide con su dirección y no encuentra cómo corregirlo. Los otros momentos críticos esperados son M2 · Explorar el mapa (−2, "
          "confusión), por la contradicción entre mapa y dirección, y M4 · Guardar/confirmar (−2, incertidumbre), porque el botón de guardar "
          "parece deshabilitado y la única forma de avanzar es «OMITIR». El promedio esperado del original es −1,8 y el del rediseño, +1,6.")

    b.sub("4. Actividad 3 · Matriz UX")
    b.par("La Tabla 3 relaciona cada momento con la emoción esperada, su valoración, la evidencia observable, la mejora propuesta y el principio "
          "de IHC que la fundamenta. La versión completa, con las hojas de verificación y heurísticas, está en matriz/matriz_ux.xlsx del repositorio.")
    b.tabla("Matriz momento / emoción / evidencia / mejora (emoción y valor esperados)",
            ["Momento", "Emoción esperada", "Valor", "Evidencia", "Mejora propuesta", "Fundamento IHC"],
            [(f"{m['id']} · {m['nombre']}\n({m['captura']})", m["emocion"], f"{m['esperado_antes']:+d}".replace("-", "−") + " → " + f"{m['esperado_despues']:+d}",
              m["evidencia"], m["mejora"], m["principio"]) for m in MOMENTOS],
            [2.0, 2.3, 1.1, 3.4, 3.2, 2.8], size=8, center_cols=(2,))

    b.sub("5. Actividad 4 · Rediseño en Figma y comparación antes/después")
    b.par([("Prototipo en Figma: ", "b"), FIGMA])
    b.par("El archivo contiene la pantalla original recreada en capas editables con los problemas anotados, la pantalla rediseñada con las "
          "mejoras anotadas (M1–M8), el estado de confirmación y un estado de alerta para cuando el pin queda lejos de la dirección. "
          "El botón «Confirmar ubicación y continuar» lleva al estado confirmado en el modo de presentación de Figma.")
    b.figura("rediseño/antes_despues.png", 15.5, "Comparación visual: original con problemas C1–C6 (izquierda) y rediseño con mejoras M1–M8 (derecha).")
    b.tabla("Modificaciones del rediseño", ["ID", "Modificación", "Resuelve"], MEJORAS, [1.1, 11.4, 2.3], center_cols=(0, 2))
    b.par("Las tres modificaciones más relevantes son:")
    b.bullet([("Claridad (M4, M6): ", "b"), "el pin aparece sobre el lote de la dirección, con la etiqueta «Tu entrega» y un indicador que dice si "
              "coincide con la dirección escrita. Se elimina la contradicción que causaba la confusión de M2."])
    b.bullet([("Control (M4, M5): ", "b"), "buscador, «Usar mi ubicación», pista de arrastre y zoom de 44 × 44 px. El usuario ve cómo corregir la "
              "ubicación, lo que ataca la frustración esperada en M3."])
    b.bullet([("Confianza y retroalimentación (M7, M8, M1): ", "b"), "un solo botón primario sustituye a «OMITIR», el guardado se confirma "
              "con un mensaje y la opción «Deshacer», y se retira la cuenta regresiva. Así se reducen la incertidumbre de M4 y la inseguridad de M5."])
    b.figura("rediseño/rediseno_estado2_confirmado.png", 12.5, "Rediseño después de confirmar: mensaje de éxito, «Deshacer» y paso 4 completado.")
    b.figura("rediseño/rediseno_estado_alerta_pin_lejos.png", 12.5, "Rediseño con el pin lejos de la dirección: aviso preventivo antes de confirmar.")

    b.sub("6. Actividad 5 · Verificación de las mejoras")
    b.par([("Alcance: ", "b"), "no se realizó una nueva prueba de usuario. La verificación se hizo por inspección sobre ambos diseños, con "
           "criterios que sí se pueden medir: contraste (WCAG 2.2), tamaño de los objetivos, elementos visibles y cumplimiento de heurísticas. "
           "Estos resultados demuestran mejoras de usabilidad y accesibilidad, pero no prueban por sí solos un cambio en la emoción de los usuarios."])
    b.tabla("Verificación del rediseño por inspección (valores medidos sobre los diseños)", ["Criterio", "Original", "Rediseño", "Referencia"],
            VERIFICACION, [5.4, 2.8, 4.0, 2.6], center_cols=(1, 2))
    b.tabla("Evaluación heurística: hallazgos del original y solución en el rediseño", ["Heurística", "Original", "Rediseño"],
            HEURISTICAS, [3.6, 5.6, 5.6])

    b.sub("7. Actividad 6 · Publicación en GitHub")
    b.par([("Repositorio del grupo: ", "b"), a.repo or "https://github.com/________/________"])
    b.par("El repositorio sigue la estructura solicitada en la guía (evidencias, curva_emocional, matriz, rediseño e informe). Cada integrante "
          "publicó su parte con commits desde su propia cuenta, como se detalla en la Tabla 7.")
    b.tabla("Roles y contribuciones de los integrantes", ["Integrante", "Rol y aporte", "Archivos a su cargo"],
            [(x[1], x[2], x[3]) for x in INTEGRANTES], [3.4, 6.6, 4.8])

    # ---------------- 2.7 Resultados obtenidos
    paras = d.paragraphs
    _, h_hab = find("Habilidades blandas")
    borrar_entre(h_res._p, h_hab._p)
    b.at(h_res._p)
    b.par("", after=0)
    b.bullet("Se definieron cinco momentos de interacción y se documentaron seis problemas con evidencia numerada (C1–C6). El más grave es que el mapa contradice la dirección escrita (C3).")
    b.bullet("La curva emocional esperada del original tiene un promedio de −1,8, con momentos críticos en M3 (−3), M2 (−2) y M4 (−2). Estos valores son una hipótesis de la inspección y no fueron medidos con usuarios.")
    b.bullet("La matriz UX relaciona los cinco momentos con su emoción esperada, evidencia, mejora y fundamento de IHC (Tabla 3).")
    b.bullet("El rediseño en Figma incluye ocho modificaciones (M1–M8) y tres estados: confirmar, confirmado y alerta.")
    b.bullet("En la verificación por inspección, el botón principal pasa de un contraste de 2,3:1 a 6,6:1, los controles de zoom de 32 × 30 px a 44 × 44 px, los elementos de urgencia de 3 a 0, y se atienden las 8 heurísticas con hallazgos (Tablas 5 y 6).")
    b.par([("¿Qué momento generó la emoción más negativa y qué evidencia lo demuestra?", "b")], before=6, after=2)
    b.par("Según la inspección, el momento con la emoción más negativa esperada es M3 · Ajustar/corregir (−3, frustración). La evidencia "
          "está en la interfaz: el pin está en otra manzana que la dirección (C3) y no hay buscador, pista de arrastre ni controles de zoom "
          "accesibles (C4). Esta conclusión es una hipótesis: no fue confirmada con usuarios.")
    b.par([("¿Cómo se relacionan las mejoras con la usabilidad, la accesibilidad y la confianza?", "b")], before=6, after=2)
    b.par("Usabilidad: el pin alineado, el buscador y el botón primario único reducen errores y pasos innecesarios (heurísticas H1, H3 y H5 "
          "de Nielsen; Ley de Fitts). Accesibilidad: todos los textos y botones del rediseño superan 4,5:1 de contraste, los controles de "
          "zoom miden 44 × 44 px y los estados se comunican con ícono y texto, no solo con color. Confianza: se retiran la urgencia artificial y "
          "los errores ortográficos, se muestran todos los costos y se confirma de forma explícita que la dirección quedó guardada, lo que "
          "corresponde al nivel reflexivo del diseño emocional [1].")
    b.par([("¿Qué diferencia se identificó entre la experiencia antes y después del rediseño?", "b")], before=6, after=2)
    b.par("La diferencia emocional no se pudo medir porque no se realizó la prueba con usuarios. Las diferencias verificables están en el "
          "diseño (Tabla 5): el rediseño elimina la contradicción entre mapa y dirección, ofrece tres formas visibles de corregir la ubicación, "
          "confirma el guardado y cumple los criterios de contraste. La curva esperada sube de −1,8 a +1,6, pero esa mejora es una hipótesis "
          "que debe validarse con participantes.")

    # ---------------- 2.8 Habilidades blandas
    paras = d.paragraphs
    for label in ("Trabajo en equipo", "Comunicación asertiva", "La empatía", "Pensamiento crítico", "Responsabilidad"):
        _, p = find(label); marcar_checkbox(p, True)

    # ---------------- 2.9 Conclusiones / 2.10 Recomendaciones / 2.11 Referencias / 2.12 Anexos
    paras = d.paragraphs
    _, h_con = find("Conclusiones"); _, h_rec = find("Recomendaciones"); _, h_ref = find("Referencias bibliográficas"); _, h_anx = find("Anexos")
    for x, y in ((h_con, h_rec), (h_rec, h_ref), (h_ref, h_anx)):
        borrar_entre(x._p, y._p)
    el = h_anx._p.getnext()
    while el is not None and el.tag != qn("w:sectPr"):
        nxt = el.getnext(); el.getparent().remove(el); el = nxt

    b.at(h_con._p)
    b.par("", after=0)
    b.par("El problema central de la pantalla B-4 no es estético sino de coherencia: el mapa contradice la dirección que el usuario debe "
          "verificar y no ofrece cómo corregirla. Por eso el momento crítico esperado es M3 · Ajustar/corregir.")
    b.par("La jerarquía visual del original empuja al usuario a saltarse la verificación: «OMITIR» es 5,4 veces más grande que «Guardar "
          "dirección», que además no cumple el contraste mínimo (2,3:1) y no da respuesta al pulsarlo. El rediseño invierte esa jerarquía y "
          "agrega retroalimentación.")
    b.par("El rediseño cumple los criterios medibles de la verificación: contraste de al menos 4,5:1 en textos y botones, controles de 44 × 44 px, "
          "ningún elemento de urgencia artificial y solución para las ocho heurísticas de Nielsen con hallazgos.")
    b.par("La principal limitación del trabajo es que el efecto emocional no fue medido: la curva es una hipótesis de la inspección y, como "
          "advierte la guía, la estética por sí sola no demuestra una mejora emocional.", after=8)

    b.at(h_rec._p)
    b.par("", after=0)
    b.bullet("Validar la curva emocional con al menos tres usuarios ajenos al equipo, con la misma tarea y la escala de −3 a +3, en el original y en el rediseño. El repositorio incluye un instrumento de prueba preparado para ese fin (evidencias/pruebas_usuarios).")
    b.bullet("En una versión funcional, ubicar el pin automáticamente a partir de la dirección escrita mediante un servicio de geocodificación.")
    b.bullet("Aplicar el mismo análisis al resto del flujo (B-1 a B-5); en el paso de pago, por ejemplo, la suscripción de $9,99 al mes aparece marcada por defecto.")
    b.bullet("Probar el rediseño con teclado y lector de pantalla cuando exista una versión funcional, para completar la verificación de accesibilidad.", after=8)

    b.at(h_ref._p)
    b.par("", after=0)
    for k, ref in enumerate(REFERENCIAS, 1):
        p = b.par(f"[{k}]\t{ref}", after=3, size=10)
        p.paragraph_format.left_indent = Cm(1.6); p.paragraph_format.first_line_indent = Cm(-1.0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    b.par("", after=4)

    b.at(h_anx._p)
    b.par("", after=0)
    b.par([("Anexo A. Capturas numeradas de la pantalla original", "b")], after=6, keep=True)
    capt = [("C1_urgencia_cuenta_regresiva.png", 15, "Urgencia artificial"), ("C2_ortografia_tipografias.png", 15, "Errores ortográficos y tipografías mezcladas"),
            ("C3_pin_no_coincide_direccion.png", 11, "El pin no coincide con la dirección"), ("C4_sin_buscador_zoom_pequeno.png", 7, "Sin buscador y zoom pequeño"),
            ("C5_boton_guardar_parece_deshabilitado.png", 8, "«Guardar dirección» parece deshabilitado"), ("C6_omitir_boton_dominante.png", 15, "«OMITIR» domina la pantalla")]
    for k, (f, w, t) in enumerate(capt, 1):
        b.figura_anexo(f"evidencias/original/capturas_numeradas/{f}", w, f"Figura A.{k}", f"C{k} · {t}.")
    b.par([("Anexo B. Pantallas de contexto del flujo original", "b")], after=6, keep=True, before=6)
    b.figura_anexo("evidencias/original/B-3_contexto_paso_anterior.jpeg", 11, "Figura B.1", "Paso anterior: B-3 «Datos de quien recibe».")
    b.figura_anexo("evidencias/original/B-5_contexto_paso_siguiente.jpeg", 11, "Figura B.2", "Paso siguiente: B-5 «Finaliza tu compra».")
    b.par([("Anexo C. Enlaces", "b")], after=4, keep=True, before=6)
    b.bullet([("Prototipo en Figma: ", "b"), FIGMA], left=True)
    b.bullet([("Repositorio del grupo: ", "b"), a.repo or "https://github.com/________/________"], left=True)
    b.bullet([("Prototipo original: ", "b"), ORIGINAL_URL], left=True)

    d.core_properties.title = "Evaluación y rediseño de la experiencia emocional del usuario"
    d.core_properties.author = "Grupo IHC · 5to Software · FISEI-UTA"
    os.makedirs(os.path.dirname(a.salida), exist_ok=True)
    d.save(a.salida)
    print("Guardado:", a.salida)


if __name__ == "__main__":
    main()
